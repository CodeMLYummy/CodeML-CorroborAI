import contextlib
import csv
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from corroborai.cli import main
from corroborai.engine import corroborate
from corroborai.io.loaders import load_bundle
from corroborai.report import write_csv, write_report
from corroborai.rules_config import load_rules
from tests._helpers import DATA_DIR, requires_data

RECALC = Path("/mnt/skills/public/xlsx/scripts/recalc.py")


@requires_data
class TestReport(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.result = corroborate(load_bundle(DATA_DIR), load_rules(), strict=True)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.xlsx = write_report(cls.result, Path(cls.tmp.name) / "r.xlsx")
        cls.csv = write_csv(cls.result, Path(cls.tmp.name) / "v.csv")
        cls.wb = load_workbook(cls.xlsx)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sheets(self):
        self.assertEqual(self.wb.sheetnames, ["Synthèse", "À investiguer", "Motifs", "Règles candidates",
                                              "Écarts justifiés", "Conformes", "Détail", "Affectations",
                                              "Interprétations", "Intégrité"])

    def test_analysis_sheets(self):
        self.assertEqual(self.wb["Motifs"].max_row - 1, len(self.result.analysis.patterns))
        self.assertEqual(self.wb["Règles candidates"].max_row - 1, len(self.result.analysis.candidate_rules))

    def test_row_counts_match_verdicts(self):
        c = self.result.counts()
        self.assertEqual(self.wb["À investiguer"].max_row - 1, c["ANOMALIE"] + c["INDETERMINE"])
        self.assertEqual(self.wb["Écarts justifiés"].max_row - 1, c["JUSTIFIE"])
        self.assertEqual(self.wb["Conformes"].max_row - 1, c["CONFORME"])
        self.assertEqual(self.wb["Détail"].max_row - 1, len(self.result.findings))
        self.assertEqual(self.wb["Affectations"].max_row - 1, len(self.result.pairs))

    def test_investigate_sheet_sorted_by_priority(self):
        ws = self.wb["À investiguer"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[0], "Priorité")
        prio = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(prio, sorted(prio, reverse=True))
        for h in ("Justification", "Cause probable", "Calcul de la priorité"):
            self.assertIn(h, headers)

    def test_detail_columns_for_formulas(self):
        ws = self.wb["Détail"]
        self.assertEqual(ws["D1"].value, "Champ cible")
        self.assertEqual(ws["E1"].value, "Verdict")

    def test_integrity_sheet(self):
        ws = self.wb["Intégrité"]
        statuses = [r[4] for r in ws.iter_rows(min_row=2, values_only=True) if r[2] == "oui"]
        self.assertTrue(statuses and all(s in ("OK", "OK_RENOMME") for s in statuses))

    def test_csv(self):
        raw = self.csv.read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(len(rows) - 1, len(self.result.findings))
        self.assertEqual(rows[0][0], "Identifiant")

    @unittest.skipUnless(shutil.which("soffice") and RECALC.is_file(), "LibreOffice indisponible")
    def test_formulas_evaluate_to_python_counts(self):
        target = Path(self.tmp.name) / "recalc.xlsx"
        shutil.copy(self.xlsx, target)
        out = subprocess.run([sys.executable, str(RECALC), str(target), "90"], capture_output=True, text=True)
        report = json.loads(out.stdout)
        self.assertEqual(report.get("status"), "success", report)
        ws = load_workbook(target, data_only=True)["Synthèse"]
        values = {r[0]: r[1] for r in ws.iter_rows(min_row=1, values_only=True) if r[0]}
        for verdict, n in self.result.counts().items():
            self.assertEqual(values[verdict], n)
        self.assertEqual(values["Total"], len(self.result.findings))


@requires_data
class TestRunCli(unittest.TestCase):
    def test_run(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                code = main(["run", "--data-dir", str(DATA_DIR), "--out", tmp])
            self.assertEqual(code, 0)
            self.assertTrue((Path(tmp) / "rapport_corroboration.xlsx").is_file())
            self.assertTrue((Path(tmp) / "verdicts.csv").is_file())
            self.assertIn("ANOMALIE", out.getvalue())

    def test_run_with_override_and_bad_override(self):
        with tempfile.TemporaryDirectory() as tmp:
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(main(["run", "--data-dir", str(DATA_DIR), "--out", tmp,
                                       "--interpretation", "INT-ASSIGN-DATES=strict_literal"]), 0)
            with contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(main(["run", "--data-dir", str(DATA_DIR), "--out", tmp,
                                       "--interpretation", "INT-ASSIGN-DATES"]), 2)


if __name__ == "__main__":
    unittest.main()
