"""Fichier de rétroaction : absent, vide, exemple livré ; messages explicites (jamais d'échec silencieux)."""

import contextlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from corroborai.cli import main
from corroborai.engine import corroborate
from corroborai.feedback import FeedbackStore, store_status
from corroborai.io.loaders import load_bundle
from corroborai.report import write_report
from corroborai.rules_config import load_rules
from tests._helpers import CHALLENGE_DIR, REPO, fixture_dir, requires_challenge_data

CFG = load_rules()
EXAMPLE = REPO / "feedback" / "exemple_retroaction.yaml"


class TestStoreStatus(unittest.TestCase):
    def test_missing_file(self):
        store = FeedbackStore.load("/inexistant/retroaction.yaml")
        self.assertFalse(store.exists)
        self.assertIn("introuvable", store_status(store))
        self.assertIn("exemple_retroaction.yaml", store_status(store))

    def test_empty_then_saved(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "r.yaml"
            store = FeedbackStore.load(p)
            store.save(p)
            self.assertTrue(store.exists)
            reloaded = FeedbackStore.load(p)
            self.assertTrue(reloaded.exists and reloaded.empty)
            self.assertIn("vide", store_status(reloaded))

    def test_non_empty(self):
        self.assertIsNone(store_status(FeedbackStore.load(EXAMPLE)))


class TestMissingFileIsReported(unittest.TestCase):
    def test_cli_warns(self):
        with tempfile.TemporaryDirectory() as tmp:
            out, err = io.StringIO(), io.StringIO()
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
                code = main(["run", "--data-dir", str(fixture_dir()), "--out", tmp,
                             "--retroaction", str(Path(tmp) / "absent.yaml")])
            self.assertEqual(code, 0)
            self.assertIn("AVERTISSEMENT", err.getvalue())
            self.assertIn("introuvable", err.getvalue())
            self.assertNotIn("Rétroaction experte :", out.getvalue())

    def test_report_explains_empty_sheet(self):
        store = FeedbackStore.load("/inexistant/retroaction.yaml")
        res = corroborate(load_bundle(fixture_dir()), CFG, strict=True, feedback=store)
        with tempfile.TemporaryDirectory() as tmp:
            wb = load_workbook(write_report(res, Path(tmp) / "r.xlsx"))
            rows = list(wb["Rétroaction experte"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), 1)
            self.assertIn("introuvable", " ".join(str(v) for v in rows[0] if v))


class TestExampleFile(unittest.TestCase):
    def test_example_is_valid_and_generic_rules_apply_elsewhere(self):
        store = FeedbackStore.load(EXAMPLE, set(CFG.targets))
        self.assertEqual(len(store.rules), 3)
        res = corroborate(load_bundle(fixture_dir()), CFG, strict=True, feedback=store)
        applied = {k: len(v) for k, v in res.feedback.rules_applied.items()}
        self.assertEqual(applied, {"R-EXPERT-001": 1, "R-EXPERT-002": 1, "R-EXPERT-003": 0})
        # La correction vise un écart des fichiers du défi : signalée, jamais appliquée ailleurs
        self.assertEqual(res.feedback.corrections_missing, ["1545850:P:35342:contactEmail"])

    @requires_challenge_data
    def test_example_on_challenge_data(self):
        store = FeedbackStore.load(EXAMPLE, set(CFG.targets))
        base = corroborate(load_bundle(CHALLENGE_DIR), CFG, strict=True)
        res = corroborate(load_bundle(CHALLENGE_DIR), CFG, strict=True, feedback=store)
        self.assertEqual({k: len(v) for k, v in res.feedback.rules_applied.items()},
                         {"R-EXPERT-001": 4, "R-EXPERT-002": 4, "R-EXPERT-003": 22})
        self.assertEqual(res.feedback.corrections_applied, ["1545850:P:35342:contactEmail"])
        self.assertEqual(res.counts()["ANOMALIE"], base.counts()["ANOMALIE"] - 31)


if __name__ == "__main__":
    unittest.main()
