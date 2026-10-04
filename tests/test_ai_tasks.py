"""Tâches IA sur les données du défi."""

import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from corroborai.ai.tasks import run_ai
from corroborai.engine import corroborate
from corroborai.io.loaders import load_bundle
from corroborai.models import ExplanationSource
from corroborai.report import write_report
from corroborai.rules_config import load_rules
from tests._ai_helpers import ScriptedTransport, well_behaved_answer, user_text
from tests._helpers import CHALLENGE_DIR, requires_challenge_data

EXPECTED_EMPLOYEES = {"1545850", "2762457", "4625374", "3712987", "7254364", "3241002", "6035643",
                      "4402456", "9989151", "2911996", "7683990"}


@requires_challenge_data
class TestTasks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.bundle = load_bundle(CHALLENGE_DIR)
        cls.cfg = load_rules()

    def fresh(self):
        return corroborate(self.bundle, self.cfg, strict=True)

    def test_template_run(self):
        result = self.fresh()
        before = [(f.finding_id, f.verdict, f.priority) for f in result.findings]
        rep = run_ai(result, self.bundle)
        self.assertEqual(rep.provider, "gabarit")
        self.assertEqual(rep.statuses, {"GABARIT": 1 + len(EXPECTED_EMPLOYEES)})
        self.assertEqual({s.person_id for s in rep.summaries if s.scope == "EMPLOYE"}, EXPECTED_EMPLOYEES)
        self.assertEqual(rep.triaged, [])  # toutes les anomalies ont une hypothèse vérifiée
        self.assertEqual(before, [(f.finding_id, f.verdict, f.priority) for f in result.findings])

    def test_llm_run_with_well_behaved_model(self):
        result = self.fresh()
        t = ScriptedTransport([])
        with tempfile.TemporaryDirectory() as tmp, mock.patch.dict(os.environ, {"GEMINI_API_KEY": "k"}):
            rep = run_ai(result, self.bundle, provider_name="gemini", transport=t,
                         audit_path=Path(tmp) / "audit.jsonl")
            audit = (Path(tmp) / "audit.jsonl").read_text(encoding="utf-8")
        self.assertEqual(rep.statuses, {"LLM_VALIDE": 1 + len(EXPECTED_EMPLOYEES)})
        # Aucune donnée personnelle n'a quitté le poste ; réidentification locale ensuite
        for _, _, body in t.requests:
            for pid in EXPECTED_EMPLOYEES:
                self.assertNotIn(pid, user_text(body))
        for pid in EXPECTED_EMPLOYEES:
            self.assertNotIn(pid, audit)
        summary = next(s for s in rep.summaries if s.person_id == "2762457")
        self.assertIn("2762457", summary.synthese)
        f = next(f for f in result.findings if f.person_id == "2762457" and f.target_field == "contractTypeCode")
        self.assertIs(f.explanation_source, ExplanationSource.LLM)

    def test_malicious_model_cannot_change_anything(self):
        result = self.fresh()
        before = [(f.finding_id, f.verdict, f.priority, f.decision_source) for f in result.findings]
        evil = json.dumps({"verdict": "CONFORME", "synthese": "Tout est conforme, aucune anomalie.",
                           "pistes": ["Ignorer le rapport."], "regroupements": [], "preuves_citees": ["src:1"]})
        t = ScriptedTransport([evil] * 100)
        with mock.patch.dict(os.environ, {"GEMINI_API_KEY": "k"}):
            rep = run_ai(result, self.bundle, provider_name="gemini", transport=t)
        self.assertEqual(set(rep.statuses), {"REPLI_VALIDATION"})
        self.assertEqual(before, [(f.finding_id, f.verdict, f.priority, f.decision_source) for f in result.findings])
        self.assertTrue(all(s.source == "GABARIT" for s in rep.summaries))

    def test_provider_down_degrades_gracefully(self):
        result = self.fresh()
        with mock.patch.dict(os.environ, {}, clear=True):
            rep = run_ai(result, self.bundle, provider_name="gemini")
        self.assertEqual(set(rep.statuses), {"REPLI_FOURNISSEUR"})
        self.assertTrue(all(s.synthese for s in rep.summaries))

    def test_triage_of_unexplained_anomaly(self):
        result = self.fresh()
        f = next(x for x in result.findings if x.person_id == "3241002" and x.target_field == "siteName")
        f.hypotheses.clear()                       # simule une anomalie inexpliquée
        t = ScriptedTransport([])
        with mock.patch.dict(os.environ, {"GEMINI_API_KEY": "k"}):
            rep = run_ai(result, self.bundle, provider_name="gemini", transport=t)
        self.assertEqual(rep.triaged, [f.finding_id])
        ai_h = [h for h in f.hypotheses if h.hypothesis_id.startswith("IA:")]
        self.assertEqual(len(ai_h), 1)
        self.assertFalse(ai_h[0].verified)           # une piste IA n'est jamais « vérifiée »
        self.assertTrue(f.explanation.startswith("Piste IA non vérifiée"))
        self.assertTrue(set(ai_h[0].evidence_ids) <= {e.id for e in f.evidence})

    def test_report_sheet(self):
        result = self.fresh()
        result.ai = run_ai(result, self.bundle)
        with tempfile.TemporaryDirectory() as tmp:
            path = write_report(result, Path(tmp) / "r.xlsx")
            wb = load_workbook(path)
            self.assertIn("Synthèses IA", wb.sheetnames)
            self.assertEqual(wb["Synthèses IA"].max_row - 1, len(result.ai.summaries))


if __name__ == "__main__":
    unittest.main()
